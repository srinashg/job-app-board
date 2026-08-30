"""Extract structured, user-editable content from an uploaded resume.

Parsing is intentionally heuristic and deterministic: everything it produces is
shown to the user for correction, and their edits are what matching reads.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date
from typing import Any

from app.services.taxonomy import canonical_skills, extract_technologies, looks_like_job_title
from app.services.text import normalize

SUPPORTED_CONTENT_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/msword": "docx",
    "text/plain": "txt",
    "text/markdown": "txt",
}

SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".md": "txt"}

_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}")
_URL_RE = re.compile(r"(?:https?://)?(?:www\.)?(linkedin\.com|github\.com)/[\w\-/.]+", re.IGNORECASE)
_YEAR_RE = re.compile(r"(?:19|20)\d{2}")
_MONTH = (
    r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)"
    r"(?:uary|ruary|ch|il|e|y|ust|tember|ober|ember)?\.?"
)
#: "Jan 2021 - Present", "2018 - 2021", "March 2019 to Dec 2020". The month
#: name is matched explicitly so a trailing word (a company name, say) is not
#: absorbed into the start date.
_DATE_RANGE_RE = re.compile(
    r"(?P<start>(?:" + _MONTH + r"\s+)?(?:19|20)\d{2})"
    r"\s*(?:-|–|—|to)\s*"
    r"(?P<end>(?:" + _MONTH + r"\s+)?(?:19|20)\d{2}|present|current|now)",
    re.IGNORECASE,
)
_YEARS_EXPERIENCE_RE = re.compile(
    r"(\d{1,2})(?:\s*\+)?\s*(?:\+\s*)?years?(?:\s+of)?\s+(?:professional\s+|relevant\s+|industry\s+)?experience",
    re.IGNORECASE,
)
_BULLET_RE = re.compile(r"^\s*[-•*·▪●o]\s+")
_SPLIT_RE = re.compile(r"[,;|/]| and | & ")

SECTION_ALIASES: dict[str, tuple[str, ...]] = {
    "summary": ("summary", "profile", "objective", "about me", "professional summary"),
    "skills": ("skills", "technical skills", "core competencies", "technologies", "toolkit"),
    "experience": (
        "experience",
        "work experience",
        "professional experience",
        "employment",
        "employment history",
        "work history",
        "relevant experience",
    ),
    "education": ("education", "academic background", "academics", "degrees"),
    "certifications": (
        "certifications",
        "certificates",
        "licenses",
        "licenses and certifications",
        "credentials",
    ),
    "projects": ("projects", "personal projects", "selected projects"),
}

_HEADING_LOOKUP: dict[str, str] = {
    normalize(alias): section
    for section, aliases in SECTION_ALIASES.items()
    for alias in aliases
}


class ResumeParseError(ValueError):
    """Raised when a resume file cannot be read at all."""


@dataclass
class ParsedResume:
    full_name: str | None = None
    email: str | None = None
    phone: str | None = None
    location: str | None = None
    summary: str | None = None
    skills: list[str] = field(default_factory=list)
    job_titles: list[str] = field(default_factory=list)
    experience: list[dict[str, Any]] = field(default_factory=list)
    education: list[dict[str, Any]] = field(default_factory=list)
    certifications: list[dict[str, Any]] = field(default_factory=list)
    years_of_experience: float | None = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "full_name": self.full_name,
            "email": self.email,
            "phone": self.phone,
            "location": self.location,
            "summary": self.summary,
            "skills": self.skills,
            "job_titles": self.job_titles,
            "experience": self.experience,
            "education": self.education,
            "certifications": self.certifications,
            "years_of_experience": self.years_of_experience,
        }


def detect_kind(content_type: str | None, filename: str | None) -> str | None:
    if content_type and content_type.split(";")[0].strip() in SUPPORTED_CONTENT_TYPES:
        return SUPPORTED_CONTENT_TYPES[content_type.split(";")[0].strip()]
    if filename:
        lowered = filename.lower()
        for extension, kind in SUPPORTED_EXTENSIONS.items():
            if lowered.endswith(extension):
                return kind
    return None


def extract_text(data: bytes, content_type: str | None = None, filename: str | None = None) -> str:
    """Pull plain text out of a PDF, DOCX or text resume."""
    kind = detect_kind(content_type, filename)
    if kind is None:
        raise ResumeParseError(
            "Unsupported resume format. Upload a PDF, DOCX or plain-text file."
        )
    if kind == "pdf":
        return _extract_pdf_text(data)
    if kind == "docx":
        return _extract_docx_text(data)
    return data.decode("utf-8", errors="replace")


def _extract_pdf_text(data: bytes) -> str:
    try:
        from pdfminer.high_level import extract_text as pdf_extract_text
    except ImportError as exc:  # pragma: no cover - dependency is pinned
        raise ResumeParseError("PDF support is not available on this server") from exc
    try:
        return pdf_extract_text(io.BytesIO(data)) or ""
    except Exception as exc:  # pdfminer raises a wide range of parse errors
        raise ResumeParseError(f"Could not read the PDF: {exc}") from exc


def _extract_docx_text(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency is pinned
        raise ResumeParseError("DOCX support is not available on this server") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeParseError(f"Could not read the DOCX file: {exc}") from exc
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _clean_lines(text: str) -> list[str]:
    lines = []
    for raw in text.splitlines():
        line = raw.replace("\xa0", " ").rstrip()
        line = re.sub(r"[ \t]{2,}", "  ", line)
        lines.append(line)
    return lines


def _heading_for(line: str) -> str | None:
    stripped = line.strip().rstrip(":").strip()
    if not stripped or len(stripped) > 60:
        return None
    if _BULLET_RE.match(line):
        return None
    return _HEADING_LOOKUP.get(normalize(stripped))


def split_sections(text: str) -> dict[str, list[str]]:
    """Split resume text into named sections, keyed by canonical section name."""
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in _clean_lines(text):
        heading = _heading_for(line)
        if heading:
            current = heading
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return sections


def _first_nonempty(lines: list[str]) -> str | None:
    for line in lines:
        if line.strip():
            return line.strip()
    return None


def _parse_name(header_lines: list[str]) -> str | None:
    for line in header_lines[:6]:
        candidate = line.strip()
        if not candidate or _EMAIL_RE.search(candidate) or _PHONE_RE.search(candidate):
            continue
        words = candidate.split()
        if not 1 < len(words) <= 5:
            continue
        if any(char.isdigit() for char in candidate):
            continue
        if looks_like_job_title(candidate):
            continue
        letters = [word for word in words if word[0].isalpha()]
        if len(letters) == len(words):
            return candidate
    return None


def _parse_location(header_lines: list[str]) -> str | None:
    pattern = re.compile(r"^[A-Za-z .'-]+,\s*(?:[A-Za-z .'-]+|[A-Z]{2})(?:\s+\d{5})?$")
    for line in header_lines[:10]:
        candidate = line.strip().strip("|").strip()
        if pattern.match(candidate) and len(candidate) <= 60:
            return candidate
    return None


def _parse_skills(section_lines: list[str], full_text: str) -> list[str]:
    raw: list[str] = []
    for line in section_lines:
        cleaned = _BULLET_RE.sub("", line).strip()
        if not cleaned:
            continue
        # "Languages: Python, Go" -> drop the category label.
        if ":" in cleaned:
            head, _, tail = cleaned.partition(":")
            if len(head.split()) <= 4:
                cleaned = tail
        raw.extend(part.strip(" .•-") for part in _SPLIT_RE.split(cleaned))
    listed = canonical_skills([item for item in raw if 1 < len(item) <= 40])
    detected = extract_technologies(full_text)
    merged = list(dict.fromkeys([*listed, *detected]))
    return merged


def _split_entries(lines: list[str]) -> list[list[str]]:
    """Group a section's lines into entries, breaking on blank lines/date ranges."""
    entries: list[list[str]] = []
    current: list[str] = []
    for line in lines:
        if not line.strip():
            if current:
                entries.append(current)
                current = []
            continue
        starts_entry = bool(_DATE_RANGE_RE.search(line)) and not _BULLET_RE.match(line)
        if starts_entry and current and any(_DATE_RANGE_RE.search(item) for item in current):
            entries.append(current)
            current = []
        current.append(line)
    if current:
        entries.append(current)
    return entries


def _parse_experience(lines: list[str]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for block in _split_entries(lines):
        header = " ".join(block[:2])
        match = _DATE_RANGE_RE.search(header)
        heading_lines = [line for line in block[:2] if line.strip()]
        title = None
        company = None
        for line in heading_lines:
            cleaned = _DATE_RANGE_RE.sub("", line).strip(" ,|-–—\t")
            parts = [part.strip() for part in re.split(r"\s*(?:\||,|—|–| at | @ )\s*", cleaned) if part.strip()]
            for part in parts:
                if title is None and looks_like_job_title(part):
                    title = part
                elif company is None and part and part != title:
                    company = part
        if title is None and company is None:
            continue
        end_value = match.group("end") if match else None
        description_lines = [
            _BULLET_RE.sub("", line).strip() for line in block[2:] if line.strip()
        ]
        entries.append(
            {
                "title": title,
                "company": company,
                "start_date": match.group("start") if match else None,
                "end_date": end_value,
                "location": None,
                "description": "\n".join(description_lines) or None,
                "is_current": bool(
                    end_value and end_value.lower() in {"present", "current", "now"}
                ),
            }
        )
    return entries


def _parse_education(lines: list[str]) -> list[dict[str, Any]]:
    degree_re = re.compile(
        r"\b(ph\.?d|doctorate|m\.?s\.?c?|m\.?b\.?a|master'?s?|b\.?s\.?c?|b\.?a|bachelor'?s?|"
        r"associate'?s?|diploma|certificate)\b",
        re.IGNORECASE,
    )
    entries: list[dict[str, Any]] = []
    for block in _split_entries(lines):
        text = " ".join(line.strip() for line in block if line.strip())
        if not text:
            continue
        degree_match = degree_re.search(text)
        years = _YEAR_RE.findall(text)
        institution = None
        for line in block:
            cleaned = line.strip()
            if cleaned and not degree_re.search(cleaned) and len(cleaned) < 120:
                institution = _DATE_RANGE_RE.sub("", cleaned).strip(" ,|-–—")
                break
        if institution is None and not degree_match:
            continue
        field_match = re.search(r"\bin\s+([A-Za-z &]+)", text)
        gpa_match = re.search(r"gpa[:\s]*([0-4](?:\.\d{1,2})?)", text, re.IGNORECASE)
        entries.append(
            {
                "institution": institution,
                "degree": degree_match.group(0) if degree_match else None,
                "field_of_study": field_match.group(1).strip() if field_match else None,
                "start_date": None,
                "end_date": years[-1] if years else None,
                "gpa": gpa_match.group(1) if gpa_match else None,
            }
        )
    return entries


def _parse_certifications(lines: list[str]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for line in lines:
        cleaned = _BULLET_RE.sub("", line).strip()
        if not cleaned or len(cleaned) > 200:
            continue
        year_match = _YEAR_RE.search(cleaned)
        issuer = None
        name = cleaned
        for separator in (" - ", " – ", " | ", ", "):
            if separator in cleaned:
                name, _, issuer = cleaned.partition(separator)
                break
        name = _YEAR_RE.sub("", name).strip(" ,|-–—()")
        if not name:
            continue
        entries.append(
            {
                "name": name,
                "issuer": (issuer or "").strip(" ,|-–—()") or None,
                "issued_date": year_match.group(0) if year_match else None,
                "expires_date": None,
                "credential_id": None,
            }
        )
    return entries


def _estimate_years_of_experience(
    text: str, experience: list[dict[str, Any]]
) -> float | None:
    stated = _YEARS_EXPERIENCE_RE.search(text)
    if stated:
        return float(stated.group(1))
    years: list[int] = []
    current_year = date.today().year
    for entry in experience:
        start = entry.get("start_date")
        end = entry.get("end_date")
        start_match = _YEAR_RE.search(start or "")
        if not start_match:
            continue
        start_year = int(start_match.group(0))
        end_match = _YEAR_RE.search(end or "")
        end_year = int(end_match.group(0)) if end_match else current_year
        if end_year >= start_year:
            years.append(min(end_year, current_year) - start_year)
    if not years:
        return None
    return float(sum(years))


def parse_resume_text(text: str) -> ParsedResume:
    """Turn raw resume text into structured, editable fields."""
    sections = split_sections(text)
    header = sections.get("header", [])
    parsed = ParsedResume()

    parsed.full_name = _parse_name(header)
    email_match = _EMAIL_RE.search(text)
    parsed.email = email_match.group(0) if email_match else None
    phone_match = _PHONE_RE.search("\n".join(header) or text)
    parsed.phone = phone_match.group(0).strip() if phone_match else None
    parsed.location = _parse_location(header)

    summary_lines = [line.strip() for line in sections.get("summary", []) if line.strip()]
    parsed.summary = " ".join(summary_lines)[:2000] or None

    parsed.skills = _parse_skills(sections.get("skills", []), text)
    parsed.experience = _parse_experience(sections.get("experience", []))
    parsed.education = _parse_education(sections.get("education", []))
    parsed.certifications = _parse_certifications(sections.get("certifications", []))

    titles = [entry["title"] for entry in parsed.experience if entry.get("title")]
    if not titles:
        headline = _first_nonempty(
            [line for line in header[:8] if looks_like_job_title(line)]
        )
        if headline:
            titles = [headline.strip()]
    parsed.job_titles = list(dict.fromkeys(titles))

    parsed.years_of_experience = _estimate_years_of_experience(text, parsed.experience)
    return parsed


def parse_resume_file(
    data: bytes, content_type: str | None = None, filename: str | None = None
) -> tuple[str, ParsedResume]:
    """Extract text and parse it, returning both for storage."""
    text = extract_text(data, content_type, filename)
    if not text.strip():
        raise ResumeParseError(
            "No text could be extracted. Scanned/image-only resumes are not supported yet."
        )
    return text, parse_resume_text(text)
